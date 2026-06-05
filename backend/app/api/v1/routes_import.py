from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.session import get_db
from app.services.csv_import_service import import_csv_into_dataset
from app.auth.security import get_current_user
from app.db.models import Dataset, User
from sqlalchemy import select
from app.services.security_service import (
    NORMAL_DATASET_EXTENSIONS,
    NORMAL_DATASET_MAX_UPLOAD_BYTES,
    audit_log,
    read_upload_limited,
    validate_basic_file_signature,
)

router = APIRouter(prefix="/datasets", tags=["Dataset Import"])


class ImportCSVRequest(BaseModel):
    csv_text: str
    delimiter: str = ","
    has_header: bool = True
    max_rows: int = Field(default=5000, ge=1, le=100000)
    drop_existing: bool = True


def _rows_to_csv_text(header: list[str], rows: list[list[Any]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    return buf.getvalue()


def _json_to_csv_text(raw: bytes) -> str:
    try:
        data = json.loads(raw.decode("utf-8"))
    except UnicodeDecodeError:
        data = json.loads(raw.decode("utf-8", errors="replace"))
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail=f"invalid JSON dataset: {exc.msg}") from None

    if isinstance(data, dict):
        for key in ("rows", "data", "items", "records"):
            candidate = data.get(key)
            if isinstance(candidate, list):
                data = candidate
                break

    if not isinstance(data, list) or not data:
        raise HTTPException(status_code=400, detail="JSON dataset must be a non-empty array or contain rows/data/items")
    if not all(isinstance(item, dict) for item in data):
        raise HTTPException(status_code=400, detail="JSON dataset rows must be objects")

    header: list[str] = []
    seen: set[str] = set()
    for row in data:
        for key in row.keys():
            k = str(key)
            if k not in seen:
                seen.add(k)
                header.append(k)
    rows = [[row.get(col, "") for col in header] for row in data]
    return _rows_to_csv_text(header, rows)


def _xlsx_to_csv_text(raw: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="XLSX import requires openpyxl to be installed",
        ) from None

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"invalid XLSX dataset: {exc}") from None

    cleaned = [[("" if cell is None else cell) for cell in row] for row in rows if row and any(cell is not None for cell in row)]
    if len(cleaned) < 2:
        raise HTTPException(status_code=400, detail="XLSX dataset must contain a header row and data rows")
    header = [str(x) if str(x).strip() else f"col_{i + 1}" for i, x in enumerate(cleaned[0])]
    return _rows_to_csv_text(header, cleaned[1:])


def _uploaded_dataset_to_csv_text(raw: bytes, filename: str | None) -> tuple[str, bool]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".json":
        return _json_to_csv_text(raw), True
    if suffix == ".xlsx":
        return _xlsx_to_csv_text(raw), True
    try:
        return raw.decode("utf-8"), False
    except UnicodeDecodeError:
        return raw.decode("utf-8", errors="replace"), False


@router.post("/{dataset_id}/import_csv")
async def import_csv(
    dataset_id: int,
    req: ImportCSVRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    ds = (
        await db.execute(select(Dataset).where(Dataset.id == dataset_id, Dataset.user_id == user.id))
    ).scalar_one_or_none()
    if not ds:
        raise HTTPException(status_code=404, detail="dataset not found")
    if len(req.csv_text.encode("utf-8")) > NORMAL_DATASET_MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"csv_text exceeds {NORMAL_DATASET_MAX_UPLOAD_BYTES} bytes")
    audit_log("dataset.import_text", user_id=user.id, dataset_id=dataset_id, max_rows=req.max_rows, drop_existing=req.drop_existing)
    return await import_csv_into_dataset(
        db,
        dataset_id,
        req.csv_text,
        delimiter=req.delimiter,
        has_header=req.has_header,
        max_rows=req.max_rows,
        drop_existing=req.drop_existing,
    )


@router.post("/{dataset_id}/import_csv_file")
async def import_csv_file(
    dataset_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    file: UploadFile = File(...),
    delimiter: str = Form(","),
    has_header: bool = Form(True),
    max_rows: int = Form(5000),
    drop_existing: bool = Form(True),
) -> dict[str, Any]:
    ds = (
        await db.execute(select(Dataset).where(Dataset.id == dataset_id, Dataset.user_id == user.id))
    ).scalar_one_or_none()
    if not ds:
        raise HTTPException(status_code=404, detail="dataset not found")
    raw = await read_upload_limited(file, max_bytes=NORMAL_DATASET_MAX_UPLOAD_BYTES)
    if not raw:
        raise HTTPException(status_code=400, detail="empty file")
    if max_rows < 1 or max_rows > 100000:
        raise HTTPException(status_code=400, detail="max_rows must be between 1 and 100000")
    validate_basic_file_signature(raw, file.filename, NORMAL_DATASET_EXTENSIONS)
    audit_log("dataset.import_file", user_id=user.id, dataset_id=dataset_id, filename=Path(file.filename or "upload").name, bytes=len(raw))
    csv_text, converted_with_header = _uploaded_dataset_to_csv_text(raw, file.filename)
    return await import_csv_into_dataset(
        db,
        dataset_id,
        csv_text,
        delimiter=delimiter,
        has_header=True if converted_with_header else has_header,
        max_rows=max_rows,
        drop_existing=drop_existing,
    )
